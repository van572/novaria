import os
import re
from typing import Dict, List
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

class DocxGenerator:
    """Generador de documentos DOCX"""
    
    def __init__(self, workspace_path: str):
        self.workspace_path = workspace_path
        self.docs_folder = os.path.join(workspace_path, "documentos")
        os.makedirs(self.docs_folder, exist_ok=True)
        print(f"📁 Carpeta de documentos: {self.docs_folder}")
    
    def crear_documento(self, titulo: str, contenido: str, autor: str = "Novaria") -> Dict:
        """Crea un documento DOCX formateado"""
        try:
            # Validar entrada
            if not titulo or not contenido:
                return {
                    "success": False,
                    "error": "Título y contenido son requeridos",
                    "mensaje": "❌ Título y contenido no pueden estar vacíos"
                }
            
            # Crear documento
            doc = Document()
            
            # Configurar título
            title = doc.add_heading(titulo, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Agregar metadata
            doc.add_paragraph(f"Autor: {autor}")
            doc.add_paragraph(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph()
            
            # Agregar contenido (soporta markdown básico)
            self._agregar_contenido_formateado(doc, contenido)
            
            # Guardar documento
            nombre_archivo = self._sanitizar_nombre(titulo)
            ruta_completa = os.path.join(self.docs_folder, f"{nombre_archivo}.docx")
            doc.save(ruta_completa)
            
            # Verificar que se guardó
            if os.path.exists(ruta_completa):
                return {
                    "success": True,
                    "ruta": ruta_completa,
                    "nombre": f"{nombre_archivo}.docx",
                    "tamaño": os.path.getsize(ruta_completa),
                    "mensaje": f"✅ Documento '{titulo}' creado exitosamente en:\n{ruta_completa}"
                }
            else:
                return {
                    "success": False,
                    "error": "No se pudo verificar la creación",
                    "mensaje": "❌ Error: No se pudo guardar el documento"
                }
                
        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "mensaje": f"❌ Error creando documento: {str(e)}"
            }
    
    def _agregar_contenido_formateado(self, doc, contenido: str):
        """Agrega contenido con formato básico markdown"""
        lineas = contenido.split('\n')
        
        for linea in lineas:
            linea = linea.strip()
            if not linea:
                doc.add_paragraph()
                continue
            
            # Detectar encabezados
            if linea.startswith('# '):
                doc.add_heading(linea[2:], level=1)
            elif linea.startswith('## '):
                doc.add_heading(linea[3:], level=2)
            elif linea.startswith('### '):
                doc.add_heading(linea[4:], level=3)
            elif linea.startswith('#### '):
                doc.add_heading(linea[5:], level=4)
            
            # Detectar listas
            elif linea.startswith('- ') or linea.startswith('* '):
                doc.add_paragraph(linea[2:], style='List Bullet')
            
            # Detectar listas numeradas
            elif re.match(r'^\d+\. ', linea):
                doc.add_paragraph(linea, style='List Number')
            
            # Texto normal
            else:
                p = doc.add_paragraph(linea)
                p.style.font.size = Pt(11)
    
    def _sanitizar_nombre(self, nombre: str) -> str:
        """Sanitiza el nombre del archivo"""
        caracteres_invalidos = '<>:"/\\|?*'
        for char in caracteres_invalidos:
            nombre = nombre.replace(char, '_')
        # Limitar longitud
        return nombre[:50]
    
    def leer_documento(self, nombre_archivo: str) -> Dict:
        """Lee un documento DOCX"""
        try:
            # Asegurar que tiene extensión
            if not nombre_archivo.endswith('.docx'):
                nombre_archivo = f"{nombre_archivo}.docx"
            
            ruta = os.path.join(self.docs_folder, nombre_archivo)
            if not os.path.exists(ruta):
                return {
                    "success": False,
                    "error": "Documento no encontrado",
                    "mensaje": f"❌ No se encontró el documento: {nombre_archivo}"
                }
            
            doc = Document(ruta)
            contenido = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    contenido.append(paragraph.text)
            
            return {
                "success": True,
                "contenido": '\n'.join(contenido),
                "ruta": ruta,
                "mensaje": f"✅ Documento leído: {nombre_archivo}"
            }
            
        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "mensaje": f"❌ Error leyendo documento: {str(e)}"
            }
    
    def listar_documentos(self) -> Dict:
        """Lista todos los documentos DOCX"""
        try:
            documentos = [f for f in os.listdir(self.docs_folder) if f.endswith('.docx')]
            if not documentos:
                return {
                    "success": True,
                    "documentos": [],
                    "mensaje": "📁 No hay documentos creados aún"
                }
            
            return {
                "success": True,
                "documentos": documentos,
                "total": len(documentos),
                "mensaje": f"📁 {len(documentos)} documentos encontrados:\n" + "\n".join(f"  - {d}" for d in documentos)
            }
        except Exception as e:
            return {
                "success": False,
                "error": str(e),
                "mensaje": f"❌ Error listando documentos: {str(e)}"
            }