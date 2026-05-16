import os
import re
from datetime import datetime
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


RUTA_BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "workspace_novaria")
RUTA_DOCUMENTOS = os.path.join(RUTA_BASE, "documentos")


class GeneradorDocumentos:

    def __init__(self):
        os.makedirs(RUTA_DOCUMENTOS, exist_ok=True)

    def crear_documento(self, titulo: str, contenido: str, autor: str = "Novaria") -> dict:
        if not titulo or not contenido:
            return {"exito": False, "mensaje": "El título y el contenido son obligatorios"}

        doc = Document()
        estilo_titulo = doc.styles["Title"]
        estilo_titulo.font.size = Pt(18)
        estilo_titulo.font.bold = True

        parrafo_titulo = doc.add_paragraph(titulo, style="Title")
        parrafo_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"Generado por Novaria | {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        doc.add_paragraph(f"Autor: {autor}")
        doc.add_paragraph("")

        self._agregar_contenido_formateado(doc, contenido)

        nombre_archivo = self._sanitizar_nombre(titulo) + ".docx"
        ruta = os.path.join(RUTA_DOCUMENTOS, nombre_archivo)
        doc.save(ruta)

        return {"exito": True, "mensaje": f"Documento creado: {nombre_archivo}", "ruta": ruta}

    def _agregar_contenido_formateado(self, doc: Document, contenido: str):
        for linea in contenido.split("\n"):
            linea = linea.strip()
            if not linea:
                continue

            if re.match(r"^### ", linea):
                doc.add_heading(linea[4:], level=3)
            elif re.match(r"^## ", linea):
                doc.add_heading(linea[3:], level=2)
            elif re.match(r"^# ", linea):
                doc.add_heading(linea[2:], level=1)
            elif re.match(r"^#### ", linea):
                doc.add_heading(linea[5:], level=4)
            elif re.match(r"^- ", linea) or re.match(r"^\* ", linea):
                doc.add_paragraph(linea[2:], style="List Bullet")
            elif re.match(r"^\d+\. ", linea):
                doc.add_paragraph(re.sub(r"^\d+\. ", "", linea), style="List Number")
            else:
                doc.add_paragraph(linea)

    def _sanitizar_nombre(self, nombre: str) -> str:
        nombre = re.sub(r'[<>:"/\\|?*]', "", nombre)
        nombre = re.sub(r"\s+", "_", nombre.strip())
        return nombre[:80]

    def leer_documento(self, nombre: str) -> dict:
        ruta = os.path.join(RUTA_DOCUMENTOS, nombre if nombre.endswith(".docx") else nombre + ".docx")
        if not os.path.exists(ruta):
            return {"exito": False, "mensaje": f"No se encontró: {nombre}"}
        try:
            doc = Document(ruta)
            texto = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return {"exito": True, "contenido": texto, "ruta": ruta}
        except Exception as e:
            return {"exito": False, "mensaje": f"Error al leer: {e}"}

    def listar_documentos(self) -> dict:
        if not os.path.exists(RUTA_DOCUMENTOS):
            return {"exito": True, "documentos": []}
        archivos = sorted(
            [f for f in os.listdir(RUTA_DOCUMENTOS) if f.endswith(".docx")]
        )
        return {"exito": True, "documentos": archivos, "ruta": RUTA_DOCUMENTOS}
